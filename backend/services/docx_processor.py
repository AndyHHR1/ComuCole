import logging
from pathlib import Path

import docx

from backend.config import TEMP_DIR

logger = logging.getLogger(__name__)


def _extraer_texto_tabla(tabla: docx.table.Table) -> str:
    """Extrae texto de una tabla de Word, evitando celdas duplicadas por merging."""
    filas: list[str] = []
    for fila in tabla.rows:
        celdas = [celda.text.strip() for celda in fila.cells]
        celdas_limpias: list[str] = []
        for celda in celdas:
            if celda and (not celdas_limpias or celda != celdas_limpias[-1]):
                celdas_limpias.append(celda)
        if celdas_limpias:
            filas.append(" | ".join(celdas_limpias))
    return "\n".join(filas)


def extraer_texto_docx(ruta_archivo: str) -> str:
    """
    Extrae texto plano de un documento .docx, incluyendo párrafos y tablas.

    Args:
        ruta_archivo: Ruta al archivo .docx.

    Returns:
        Texto crudo del documento.
    """
    doc = docx.Document(ruta_archivo)
    partes: list[str] = []

    for parrafo in doc.paragraphs:
        texto = parrafo.text.strip()
        if texto:
            partes.append(texto)

    for tabla in doc.tables:
        texto_tabla = _extraer_texto_tabla(tabla)
        if texto_tabla:
            partes.append(texto_tabla)

    texto_final = "\n".join(partes)
    logger.info(
        "Texto extraido del docx (%d caracteres, %d parrafos, %d tablas)",
        len(texto_final),
        len(doc.paragraphs),
        len(doc.tables),
    )
    return texto_final
